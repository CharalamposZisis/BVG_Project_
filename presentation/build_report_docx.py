"""Generates a Word document (Prototype Development section) for the final
group report. Plain, factual content only -- no numbers or claims beyond what
was actually run and observed (see the evaluate.py output this is based on).

Usage:
    python build_report_docx.py
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_PATH = Path(__file__).resolve().parent / "Prototype_Development_Report.docx"


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Prototype Development", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    add_para(
        doc,
        "This section documents the prototype built to extract the title and date "
        "from scanned BVG construction plans, as part of the Prototype Development "
        "role for this project.",
    )

    add_heading(doc, "Approach", level=1)
    add_para(
        doc,
        "The prototype extracts title and date using zero-shot prompting of a "
        "multimodal LLM, accessed via the HTW Berlin inference API (Ollama "
        "server). No model weights are trained or modified in this approach; "
        "the model is only prompted, not fine-tuned."
    )
    add_para(
        doc,
        "The prompt used is domain-specific: it tells the model that the input "
        "is a German-language BVG construction plan, requires the title and "
        "date to be transcribed exactly as written (without translation), and "
        "explicitly instructs the model to return an empty date field rather "
        "than guess one when no date is present on the sheet."
    )

    add_heading(doc, "Implementation", level=1)
    add_para(doc, "The following components were built:")
    add_bullets(doc, [
        "evaluate.py — runs the prompt over a held-out test set and scores each "
        "prediction against manually verified ground truth: a title similarity "
        "score, a date-match classification (correct, correct-but-reformatted, "
        "or wrong), and a confusion matrix specifically for whether a date was "
        "correctly identified as present or absent (true/false positive/negative). "
        "Results are written to a CSV file for review.",
        "streamlit_app.py — a demo interface with two views: single-image upload "
        "with live extraction, and a gallery of test-set predictions showing a "
        "pass/fail indicator per image.",
        "quick_test.py — a minimal script to verify API connectivity for a "
        "single image, independent of the rest of the pipeline.",
        "analyze_failures.py — cross-references which test-set images the "
        "prototype got wrong against confidence notes recorded during ground-truth "
        "labeling, to check whether failures correlate with scans that were "
        "already hard to read.",
    ])

    add_heading(doc, "Test Set", level=1)
    add_para(
        doc,
        "The test set consists of 9 plans (18 images, counting both the "
        "full-page view and a crop of the title block for each plan), held out "
        "from the labeled dataset and not used to prepare any training data. "
        "Six of the nine plans have no date on the sheet; these were "
        "deliberately included so that the evaluation can measure whether the "
        "model hallucinates a date on plans that do not have one, not only "
        "whether it finds dates that do exist."
    )

    add_heading(doc, "Results (initial run)", level=1)
    add_para(
        doc,
        "The following results are from the first complete evaluation run "
        "against the 18-image test set:"
    )
    add_bullets(doc, [
        "Mean title similarity: 0.65",
        "Date-presence confusion matrix: 4 true positives, 2 false negatives, "
        "1 false positive, 11 true negatives",
        "Date-presence precision: 0.80, recall: 0.67",
    ])
    add_para(
        doc,
        "Two of the API calls in this run failed due to a network request "
        "timeout rather than a model error; these are recorded as empty "
        "predictions and are noted separately from genuine extraction failures.",
        italic=True,
    )

    add_heading(doc, "Identified Failure Cases", level=1)
    add_para(
        doc,
        "Two specific failure cases were identified in the initial run:"
    )
    add_bullets(doc, [
        "On one plan (S_112_002, title-block crop), the model produced an "
        "unrelated title and additionally generated a date on a sheet that has "
        "none — a confirmed hallucination. The same plan's full-page image "
        "produced a more accurate title, suggesting the crop region may not "
        "have been well aligned with the actual title block for this sheet.",
        "On another plan (1149-0007, title-block crop), the predicted title "
        "was unrelated to the ground truth, again with the full-page variant "
        "performing better than the crop.",
    ])

    add_heading(doc, "Fine-Tuning (Not Executed)", level=1)
    add_para(
        doc,
        "A LoRA fine-tuning pipeline for the same task was also implemented as "
        "an alternative approach: dataset preparation with crop and "
        "augmentation, and a training configuration for Qwen2-VL-7B-Instruct "
        "using LLaMA-Factory. This was not executed. Fine-tuning requires an "
        "NVIDIA GPU with CUDA support, which was not available on the laptops "
        "used for this project, and the HTW inference API used for the "
        "zero-shot prototype is inference-only and cannot run a training job. "
        "The fine-tuning code is included in the project repository but was "
        "not run."
    )

    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    build()
